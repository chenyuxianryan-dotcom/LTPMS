# -*- coding: utf-8 -*-
"""生成 LTPMS 导入用模拟数据 Excel（表头与后端解析严格对齐）"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import os

OUT = os.path.dirname(os.path.abspath(__file__))

HEADER_FILL = PatternFill("solid", fgColor="DBEAFE")
HEADER_FONT = Font(bold=True, size=11)


def build(filename, headers, rows, widths=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        ws.append(row)
    for idx, width in enumerate(widths or [16] * len(headers), start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    path = os.path.join(OUT, filename)
    wb.save(path)
    print(f"OK  {filename}  ({len(rows)} rows)")


# 1. 教师（工号 | 姓名 | 院系 | 联系电话）——导入后默认密码 = Ltpms@{工号}
build(
    "01_教师账号导入.xlsx",
    ["工号", "姓名", "院系", "联系电话"],
    [
        ["T001", "张伟", "计算机学院", "13800000001"],
        ["T002", "李娜", "软件学院", "13800000002"],
        ["T003", "王强", "计算机学院", "13800000003"],
        ["T004", "赵敏", "人工智能学院", "13800000004"],
        ["T005", "陈晨", "软件学院", "13800000005"],
    ],
    [12, 12, 18, 16],
)

# 2. 实验室（实验室编号 | 实验室名称 | 负责人 | 设备数）
build(
    "02_实验室导入.xlsx",
    ["实验室编号", "实验室名称", "负责人", "设备数"],
    [
        ["36-401", "网络综合实验室", "王强", 40],
        ["36-402", "软件工程实验室", "李娜", 50],
        ["37-201", "人工智能实验室", "赵敏", 30],
        ["37-202", "大数据实验室", "陈晨", 45],
        ["38-105", "嵌入式实验室", "张伟", 35],
        ["38-106", "物联网实验室", "张伟", 25],
    ],
    [14, 20, 12, 10],
)

# 3. 课表（课程号|课程名|总学时|教师工号|实验室编号|上课班级|选课人数|上课时间）
#    导入时需选择学期（先建学期）；教师工号/实验室编号须先导入
build(
    "03_课表导入.xlsx",
    ["课程号", "课程名", "总学时", "教师工号", "实验室编号", "上课班级", "选课人数", "上课时间"],
    [
        ["CS101", "数据结构实验", 32, "T001", "36-401", "计算机2401,计算机2402", 45, "周一第1-4节(1-16周)"],
        ["CS102", "操作系统实验", 24, "T002", "36-402", "软件2301", 40, "周二第1-4节(1-16周)"],
        ["AI201", "机器学习实验", 32, "T004", "37-201", "人工智能2401", 38, "周三第5-8节(1-16周)"],
        ["DB202", "数据库原理实验", 24, "T003", "37-202", "数据科学2301", 42, "周四第1-4节(1-16周)"],
        ["IOT301", "嵌入式系统实验", 28, "T005", "38-105", "物联网2201", 36, "周五第5-8节(1-16周)"],
        ["NET302", "计算机网络实验", 24, "T001", "36-401", "网络工程2301", 44, "周一第5-8节(1-16周)"],
    ],
    [10, 16, 8, 10, 12, 24, 10, 22],
)

# 4. 实验项目示例（实验名称|实验类别|实验类型|所属学科|实验要求|实验者类别|每组人数|学时数）
#    字典值与 database/dict_data.sql 一致；导入时需绑定课程（课程由课表导入自动生成）
build(
    "04_实验项目导入_CS101数据结构.xlsx",
    ["实验名称", "实验类别", "实验类型", "所属学科", "实验要求", "实验者类别", "每组人数", "学时数"],
    [
        ["线性表的基本操作", "basic", "verification", "cs", "required", "undergraduate", 2, 4],
        ["栈与队列的应用", "basic", "verification", "cs", "required", "undergraduate", 2, 4],
        ["二叉树遍历", "basic", "verification", "cs", "required", "undergraduate", 2, 4],
        ["图的遍历算法", "comprehensive", "verification", "cs", "required", "undergraduate", 2, 4],
        ["查找算法比较", "comprehensive", "design", "cs", "required", "undergraduate", 2, 4],
        ["排序算法性能分析", "comprehensive", "design", "cs", "elective", "undergraduate", 2, 4],
        ["哈夫曼编码设计", "design", "design", "cs", "elective", "undergraduate", 2, 4],
        ["综合设计:学生成绩管理系统", "comprehensive", "comprehensive", "cs", "required", "undergraduate", 4, 4],
    ],
    [26, 14, 14, 12, 12, 12, 10, 8],
)

print("ALL DONE ->", OUT)

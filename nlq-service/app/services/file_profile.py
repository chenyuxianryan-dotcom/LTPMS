"""实验材料文件画像（T07）：只读采样数据集文件，产出结构化画像统计。

红线：
  - 画像统计全部 Python 侧算好，模型只引用不计算；
  - 只读，拼接落盘路径前做 realpath 前缀校验，拒绝目录穿越（`../` / 绝对路径越界）；
  - 超限 / 坏文件 / 不支持类型一律降级为「仅元数据 + 可解释原因」，绝不抛异常。
支持类型：csv / txt 标准库采样；xlsx 经 openpyxl（read_only）同款限量；zip / 图片仅登记。
"""

import csv
import statistics
from pathlib import Path

from app.config import settings

_HEAD_ROWS = 5


def _max_bytes() -> int:
    return int(settings.profile_max_file_mb) * 1024 * 1024


def _sample_rows() -> int:
    return int(settings.profile_sample_rows)


def resolve_upload_path(stored_path: str) -> Path | None:
    """按 stored_path（UUID 文件名相对路径）拼接落盘目录并做 realpath 前缀校验。

    越界（`../` 逃逸、绝对路径覆盖）或空值返回 None —— 拒绝目录穿越。
    """
    if not stored_path:
        return None
    base = Path(settings.lab_upload_dir).resolve()
    target = (base / stored_path).resolve()
    try:
        target.relative_to(base)
    except ValueError:
        return None
    return target


def _to_number(value: str):
    text = value.strip()
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _outliers(values: list[float]) -> list[float]:
    """IQR 1.5 倍离群点（样本 <4 不判），最多回 10 个，全部 Python 侧算好。"""
    if len(values) < 4:
        return []
    ordered = sorted(values)
    q1, _, q3 = statistics.quantiles(ordered, n=4)
    iqr = q3 - q1
    low, high = q1 - 1.5 * iqr, q3 + 1.5 * iqr
    return [round(v, 4) for v in values if v < low or v > high][:10]


def _analyze_table(fields: list[str], rows: list[list[str]], truncated: bool, kind: str) -> dict:
    """csv / xlsx 共用的表格分析：字段、缺失、数值列分布与离群、样例行。"""
    missing = {name: 0 for name in fields}
    non_numeric = {name: 0 for name in fields}
    numeric_values: dict[str, list[float]] = {name: [] for name in fields}
    for row in rows:
        for idx, name in enumerate(fields):
            raw = row[idx] if idx < len(row) else ""
            text = ("" if raw is None else str(raw)).strip()
            if text == "":
                missing[name] += 1
                continue
            num = _to_number(text)
            if num is None:
                non_numeric[name] += 1
            else:
                numeric_values[name].append(num)

    numeric_columns = []
    for name in fields:
        values = numeric_values[name]
        if values and non_numeric[name] == 0:
            numeric_columns.append({
                "field": name,
                "count": len(values),
                "min": round(min(values), 4),
                "max": round(max(values), 4),
                "mean": round(statistics.fmean(values), 4),
                "median": round(statistics.median(values), 4),
                "std": round(statistics.pstdev(values), 4) if len(values) > 1 else 0.0,
                "outliers": _outliers(values),
            })

    return {
        "kind": kind,
        "row_count": len(rows),
        "column_count": len(fields),
        "fields": fields,
        "missing_by_field": missing,
        "numeric_columns": numeric_columns,
        "sample_head": rows[:_HEAD_ROWS],
        "truncated": truncated,
    }


def _profile_csv(path: Path) -> dict:
    with open(path, "r", encoding="utf-8", errors="replace", newline="") as fh:
        reader = csv.reader(fh)
        fields: list[str] = []
        rows: list[list[str]] = []
        truncated = False
        for index, row in enumerate(reader):
            if index == 0:
                fields = row
                continue
            if len(rows) >= _sample_rows():
                truncated = True
                break
            rows.append(row)
    return _analyze_table(fields, rows, truncated, "csv")


def _profile_txt(path: Path) -> dict:
    lines: list[str] = []
    total = 0
    chars = 0
    with open(path, "r", encoding="utf-8", errors="replace") as fh:
        for line in fh:
            total += 1
            chars += len(line)
            if len(lines) < _HEAD_ROWS:
                lines.append(line.rstrip("\n"))
            if total >= _sample_rows():
                break
    return {
        "kind": "txt",
        "line_count": total,
        "char_count": chars,
        "sample_head": lines,
        "truncated": total >= _sample_rows(),
    }


def _profile_xlsx(path: Path) -> dict:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_names = workbook.sheetnames
        if not workbook.worksheets:
            return {"kind": "xlsx", "sheet_names": sheet_names, "row_count": 0,
                    "column_count": 0, "fields": [], "sample_head": [], "truncated": False}
        sheet = workbook.worksheets[0]
        fields: list[str] = []
        rows: list[list[str]] = []
        truncated = False
        limit = _sample_rows()
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            cells = ["" if cell is None else str(cell) for cell in row]
            if index == 0:
                fields = cells
                continue
            if len(rows) >= limit:
                truncated = True
                break
            rows.append(cells)
        result = _analyze_table(fields, rows, truncated, "xlsx")
        result["sheet_names"] = sheet_names
        result["active_sheet"] = sheet.title
        return result
    finally:
        workbook.close()


def profile_dataset_file(stored_path: str, file_type: str,
                         original_filename: str | None = None,
                         file_size: int | None = None) -> dict:
    """对单个数据集版本文件产出画像。任何异常 / 超限 / 越界均降级为元数据，不抛出。

    返回 dict 恒含 filename / file_type / file_size / profiled；
    profiled=True 时附画像字段，False 时附 reason（可解释）。
    """
    meta = {
        "filename": original_filename or stored_path,
        "file_type": file_type,
        "file_size": file_size,
    }
    path = resolve_upload_path(stored_path)
    if path is None:
        return {**meta, "profiled": False, "reason": "路径非法或越界，已拒绝访问"}
    if not path.is_file():
        return {**meta, "profiled": False, "reason": "文件不存在或已清理"}
    try:
        actual_size = path.stat().st_size
    except OSError:
        return {**meta, "profiled": False, "reason": "文件不可读"}
    if file_size is None:
        meta["file_size"] = actual_size
    if actual_size > _max_bytes():
        return {**meta, "profiled": False,
                "reason": f"文件超过 {settings.profile_max_file_mb}MB 采样上限，仅登记元数据"}

    kind = (file_type or "").lower()
    try:
        if kind == "csv":
            return {**meta, "profiled": True, **_profile_csv(path)}
        if kind == "txt":
            return {**meta, "profiled": True, **_profile_txt(path)}
        if kind == "xlsx":
            return {**meta, "profiled": True, **_profile_xlsx(path)}
        return {**meta, "profiled": False,
                "reason": f"{kind or '未知'} 类型仅登记文件名/大小/类型，不解析内容"}
    except Exception as exc:  # 坏文件 / 依赖缺失 / 解析异常统一降级
        return {**meta, "profiled": False, "reason": f"解析失败已降级：{type(exc).__name__}"}
